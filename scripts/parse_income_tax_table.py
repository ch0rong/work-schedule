"""국세청 근로소득 간이세액표(xlsx)를 앱이 불러오는 JSON(tax_table.json)으로 바꾼다.

사용법 (프로젝트 폴더에서):
    .venv/Scripts/python scripts/parse_income_tax_table.py "근로소득 간이세액표_2026.03.01.xlsx"

출력: src/lib/tax_table.json (직접 고치지 말고 이 스크립트로 다시 만든다)

엑셀 구성 (소득세법 시행령 별표2)
- '근로소득간이세액표' 시트
  - 5행~: 월급여액(천원) [이상, 미만) 구간별 공제대상가족 1~11명 세액(원). '-'는 0원
  - '10,000천원' 행: 월급여 10,000천원일 때 세액
  - '10,000천원 초과' 이후: 구간별 계산식 문장
- '소득령 별표2' 시트: 8세 이상 20세 이하 자녀 공제액 등 부가 규칙
"""

from __future__ import annotations

import json
import re
import sys
from dataclasses import dataclass
from pathlib import Path

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

TABLE_SHEET = '근로소득간이세액표'
RULE_SHEET = '소득령 별표2'
FAMILY_COLUMNS = 11  # 공제대상가족 1~11명 (C~M열)
OUTPUT = Path(__file__).resolve().parent.parent / 'src' / 'lib' / 'tax_table.json'


@dataclass(frozen=True)
class Row:
    min_k: int  # 이상 (천원)
    max_k: int  # 미만 (천원)
    taxes: list[int]  # 가족 1~11명


@dataclass(frozen=True)
class Bracket:
    """월급여 10,000천원 초과 구간: 세액 = 10,000천원 세액 + fixed + (월급여 - over) × factor × rate"""

    over_k: int  # 이 금액(천원) 초과
    up_to_k: int | None  # 이 금액(천원) 이하 (마지막 구간은 None)
    fixed: int  # 원
    factor_percent: int  # 초과 금액에 곱하는 비율 (98% 또는 100%)
    rate_percent: int  # 세율


def to_int(value: object, where: str) -> int:
    """셀 값을 정수로. '-'는 0원, '10,000천원' 같은 문자는 숫자만 읽는다."""
    if value == '-':
        return 0
    if isinstance(value, (int, float)) and float(value).is_integer():
        return int(value)
    if isinstance(value, str) and re.fullmatch(r'[\d,]+(천원)?', value.strip()):
        return int(re.sub(r'\D', '', value))
    raise ValueError(f'{where}: 숫자가 아닌 값 {value!r}')


def parse_rows(ws: Worksheet) -> tuple[list[Row], list[int], int]:
    """본표 구간과 10,000천원 행. 반환: (구간, 10,000천원 세액, 다음 행 번호)"""
    rows: list[Row] = []
    for r in range(5, ws.max_row + 1):
        first, second = ws.cell(r, 1).value, ws.cell(r, 2).value
        taxes = [ws.cell(r, 3 + i).value for i in range(FAMILY_COLUMNS)]
        if isinstance(first, (int, float)) and isinstance(second, (int, float)):
            rows.append(Row(to_int(first, f'A{r}'), to_int(second, f'B{r}'), [to_int(v, f'{r}행') for v in taxes]))
            continue
        if isinstance(first, str) and first.replace(' ', '') == '10,000천원':
            return rows, [to_int(v, f'{r}행') for v in taxes], r + 1
        raise ValueError(f'{r}행: 예상하지 못한 형식 {first!r}')
    raise ValueError("'10,000천원' 행을 찾지 못했습니다.")


def parse_brackets(ws: Worksheet, start: int) -> list[Bracket]:
    """'N천원 초과' 행의 계산식 문장과 다음 'M천원 이하' 행에서 구간 계산식을 읽는다."""
    brackets: list[Bracket] = []
    r = start
    while r <= ws.max_row:
        label = str(ws.cell(r, 1).value or '').strip()
        over = re.fullmatch(r'([\d,]+)천원\s*초과', label)
        if not over:
            r += 1
            continue
        formula = str(ws.cell(r, 3).value or '')
        rate = re.search(r'(\d+)%\s*상당액', formula)
        if not formula.startswith('(10,000천원인 경우의 해당 세액)') or not rate:
            raise ValueError(f'{r}행: 계산식을 읽을 수 없습니다: {formula!r}')
        factor = re.search(r'(\d+)%를 곱한', formula)
        # 괄호 안의 고정 금액(원)은 모두 더한다: (1,397,000원), (25,000원)
        fixed = sum(int(amount.replace(',', '')) for amount in re.findall(r'\(([\d,]+)원\)', formula))
        upper_label = str(ws.cell(r + 1, 1).value or '').strip()
        upper = re.fullmatch(r'([\d,]+)천원\s*이하', upper_label)
        brackets.append(
            Bracket(
                over_k=int(over.group(1).replace(',', '')),
                up_to_k=int(upper.group(1).replace(',', '')) if upper else None,
                fixed=fixed,
                factor_percent=int(factor.group(1)) if factor else 100,
                rate_percent=int(rate.group(1)),
            )
        )
        r += 2 if upper else 1
    return brackets


def parse_child_credits(ws: Worksheet) -> dict[str, int]:
    """8세 이상 20세 이하 자녀 공제액 (1명, 2명, 3명 이상 추가 1명당)"""
    text = '\n'.join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
    one = re.search(r'자녀가 1명인 경우\s*:\s*([\d,]+)원', text)
    two = re.search(r'자녀가 2명인 경우\s*:\s*([\d,]+)원', text)
    extra = re.search(r'2명 초과 자녀 1명당\s*([\d,]+)원', text)
    if not (one and two and extra):
        raise ValueError('자녀 공제액을 찾지 못했습니다.')
    return {key: int(m.group(1).replace(',', '')) for key, m in (('one', one), ('two', two), ('perExtra', extra))}


def validate(rows: list[Row], at_10000: list[int], brackets: list[Bracket]) -> None:
    """구간이 빈틈없이 이어지고, 세액이 월급여에 따라 줄지 않으며, 가족이 늘면 늘지 않는지 확인한다."""
    if not rows:
        raise ValueError('본표 구간이 없습니다.')
    for prev, cur in zip(rows, rows[1:]):
        if cur.min_k != prev.max_k:
            raise ValueError(f'구간이 이어지지 않습니다: {prev.max_k} → {cur.min_k}')
        if any(c < p for c, p in zip(cur.taxes, prev.taxes)):
            raise ValueError(f'{cur.min_k}천원 구간 세액이 이전 구간보다 작습니다.')
    if rows[-1].max_k != 10000:
        raise ValueError(f'본표 마지막 구간이 10,000천원에서 끝나지 않습니다: {rows[-1].max_k}')
    for row in [*rows, Row(10000, 10000, at_10000)]:
        if len(row.taxes) != FAMILY_COLUMNS or any(b > a for a, b in zip(row.taxes, row.taxes[1:])):
            raise ValueError(f'{row.min_k}천원: 가족 수가 늘수록 세액이 줄어야 합니다.')
    for prev, cur in zip(brackets, brackets[1:]):
        if prev.up_to_k != cur.over_k:
            raise ValueError(f'초과 구간이 이어지지 않습니다: {prev.up_to_k} → {cur.over_k}')
        # 다음 구간의 고정 금액 = 이전 구간 고정 금액 + 이전 구간 전체에 대한 세액 (계산식이 끊김 없이 이어지는지)
        expected = prev.fixed + (cur.over_k - prev.over_k) * 1000 * prev.factor_percent * prev.rate_percent // 10000
        if cur.fixed != expected:
            raise ValueError(f'{cur.over_k}천원 초과 구간 고정 금액 {cur.fixed:,}원이 계산값 {expected:,}원과 다릅니다.')
    if not brackets or brackets[0].over_k != 10000 or brackets[-1].up_to_k is not None:
        raise ValueError('10,000천원 초과 구간 구성이 예상과 다릅니다.')


def render(source: Path, rows: list[Row], at_10000: list[int], brackets: list[Bracket], child: dict[str, int]) -> str:
    """앱이 불러오는 tax_table.json. 금액 단위는 모두 원(엑셀의 천원 구간은 원으로 바꾼다)."""
    effective = re.search(r'(\d{4})\.(\d{2})\.(\d{2})', source.name)
    data = {
        'source': '국세청 근로소득 간이세액표 (소득세법 시행령 제189조제1항 별표2)',
        'sourceFile': source.name,
        'effectiveFrom': '-'.join(effective.groups()) if effective else '',
        'note': '월급여액은 비과세 및 학자금 제외 금액. tax[i]는 공제대상가족 i+1명의 월 소득세(원)',
        # 본표: [min, max) 원
        'rows': [{'min': row.min_k * 1000, 'max': row.max_k * 1000, 'tax': row.taxes} for row in rows],
        # 월급여 10,000,000원일 때 세액
        'at10000k': at_10000,
        # 10,000,000원 초과: 세액 = at10000k + fixed + (월급여 - over) × factorPercent% × ratePercent%
        'brackets': [
            {'over': b.over_k * 1000, 'upTo': b.up_to_k * 1000 if b.up_to_k is not None else None, 'fixed': b.fixed,
             'factorPercent': b.factor_percent, 'ratePercent': b.rate_percent}
            for b in brackets
        ],
        'childCredit': child,
    }
    # 배열 원소(구간·계산식)는 한 줄에 하나씩 두어 사람이 읽고 비교하기 쉽게 한다
    def compact(value: object) -> str:
        return json.dumps(value, ensure_ascii=False, separators=(', ', ': '))

    lines = ['{']
    keys = list(data)
    for index, key in enumerate(keys):
        value = data[key]
        comma = ',' if index < len(keys) - 1 else ''
        if key in ('rows', 'brackets'):
            items = ',\n'.join(f'    {compact(item)}' for item in value)  # type: ignore[union-attr]
            lines.append(f'  "{key}": [\n{items}\n  ]{comma}')
        else:
            lines.append(f'  "{key}": {compact(value)}{comma}')
    lines.append('}')
    return '\n'.join(lines) + '\n'


def main() -> None:
    if len(sys.argv) != 2:
        sys.exit('사용법: python scripts/parse_income_tax_table.py <간이세액표.xlsx>')
    source = Path(sys.argv[1])
    workbook = openpyxl.load_workbook(source, data_only=True, read_only=False)
    table = workbook[TABLE_SHEET]
    rows, at_10000, next_row = parse_rows(table)
    brackets = parse_brackets(table, next_row)
    child = parse_child_credits(workbook[RULE_SHEET])
    validate(rows, at_10000, brackets)
    OUTPUT.write_text(render(source, rows, at_10000, brackets, child), encoding='utf-8', newline='\n')
    print(f'{OUTPUT.name}: 구간 {len(rows)}개({rows[0].min_k}~{rows[-1].max_k}천원), 초과 구간 {len(brackets)}개, 자녀 공제 {child}')


if __name__ == '__main__':
    main()
